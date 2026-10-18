"""Render a draft (job + document details + sections of blocks) into the real Gamsberg SOP Word template.

Strategy
- Cover block, header, footer and approval tables are Word content controls (w:sdt) tagged by
  VZI's document control system (e.g. VZI_DocNo, txtStatus). We fill them by tag and remove their
  SharePoint data bindings so Word keeps our values.
- The template body between "Purpose Scope and Users" and "Validity and Document Management" is
  example content. We remove it and insert the user's sections in the user's order.
- Headings and numbered steps reuse the template's own numbered paragraphs (list numId 8), so Word
  numbers them 1., 1.1, 1.1.1 exactly like real Gamsberg SOPs.
- Definitions and Annexures have fixed slots in the template and are filled in place.
"""
from __future__ import annotations

import copy
import datetime as dt
import io
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

TABLE_STYLE = "Grid Table 4 Accent 6"


# ------------------------------------------------------------------ helpers

def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _set_par_text(p_el, text: str, bold: bool | None = None):
    """Replace the visible text of a paragraph element, keeping its paragraph properties."""
    for child in list(p_el):
        if child.tag != qn("w:pPr"):
            p_el.remove(child)
    r = OxmlElement("w:r")
    if bold:
        rpr = OxmlElement("w:rPr")
        rpr.append(OxmlElement("w:b"))
        r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    p_el.append(r)
    return p_el


def _set_ilvl(p_el, level: int):
    ilvl = p_el.find(".//" + qn("w:ilvl"))
    if ilvl is not None:
        ilvl.set(qn("w:val"), str(level))


def _fill_sdts(root, values: dict[str, str]):
    for sdt in root.iter(qn("w:sdt")):
        pr = sdt.find(qn("w:sdtPr"))
        if pr is None:
            continue
        tag = pr.find(qn("w:tag"))
        alias = pr.find(qn("w:alias"))
        key = (tag.get(qn("w:val")) if tag is not None else "") or (alias.get(qn("w:val")) if alias is not None else "")
        if key not in values:
            continue
        for name in ("w:dataBinding", "w:showingPlcHdr"):
            for el in pr.findall(qn(name)):
                pr.remove(el)
        content = sdt.find(qn("w:sdtContent"))
        if content is None:
            continue
        ts = list(content.iter(qn("w:t")))
        if ts:
            ts[0].text = values[key]
            ts[0].set(qn("xml:space"), "preserve")
            for t in ts[1:]:
                t.text = ""
        # remove placeholder styling (grey text) if present
        for rs in content.iter(qn("w:rStyle")):
            if "Placeholder" in (rs.get(qn("w:val")) or ""):
                rs.getparent().remove(rs)


def _body_index(body, predicate):
    for i, el in enumerate(body):
        if predicate(el):
            return i
    return -1


def _is_h1(el, text: str) -> bool:
    if el.tag != qn("w:p"):
        return False
    st = el.find(".//" + qn("w:pStyle"))
    return st is not None and st.get(qn("w:val")) == "Heading1" and _text(el).strip().lower().startswith(text.lower())


# ------------------------------------------------------------------ builder

class BodyWriter:
    """Inserts elements before an anchor element in the document body."""

    def __init__(self, doc, anchor, protos):
        self.doc = doc
        self.anchor = anchor
        self.p = protos  # prototype elements cloned from the template

    def _put(self, el):
        self.anchor.addprevious(el)
        return el

    def heading1(self, text):
        self._put(_set_par_text(copy.deepcopy(self.p["h1"]), text))

    def heading2(self, text):
        self._put(_set_par_text(copy.deepcopy(self.p["h2"]), text))

    def para(self, text, bold=False, color=None):
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        self._put(p._p)

    def step(self, text, level):
        el = _set_par_text(copy.deepcopy(self.p["step"]), text)
        _set_ilvl(el, level)
        self._put(el)

    def bullet(self, text):
        p = self.doc.add_paragraph(text, style="List Bullet")
        self._put(p._p)

    def spacer(self):
        self._put(self.doc.add_paragraph()._p)

    def table(self, columns, rows):
        t = self.doc.add_table(rows=1 + len(rows), cols=len(columns))
        try:
            t.style = self.doc.styles[TABLE_STYLE]
        except KeyError:
            t.style = self.doc.styles["Table Grid"]
        for j, c in enumerate(columns):
            cell = t.rows[0].cells[j]
            cell.text = c
            for r in cell.paragraphs[0].runs:
                r.bold = True
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row[: len(columns)]):
                t.rows[i].cells[j].text = val
        for row in t.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    for r in par.runs:
                        r.font.size = Pt(9)
        self._put(t._tbl)

    def raw(self, el):
        self._put(el)

    # ---- blocks
    def blocks(self, blocks):
        has_sub = False
        for b in blocks:
            t = b["type"]
            if t == "subheading":
                self.heading2(b["text"])
                has_sub = True
            elif t == "paragraph":
                self.para(b["text"])
            elif t == "warning":
                self.para(b["text"], bold=True, color="C00000")
            elif t == "bullets":
                for it in b["items"]:
                    self.bullet(it)
            elif t == "steps":
                for it in b["items"]:
                    self.step(it, 2 if has_sub else 1)
            elif t == "table":
                self.table(b["columns"], b["rows"])
                self.spacer()


# ------------------------------------------------------------------ main

def render(template_path: Path, template_cfg: dict, job: dict, meta: dict, sections: list[dict]) -> bytes:
    doc = docx.Document(str(template_path))
    body = doc.element.body
    defaults = template_cfg.get("defaults", {})
    today = dt.date.today()

    def m(key, fallback=""):
        v = (meta or {}).get(key)
        return v if v not in (None, "") else defaults.get(key, fallback)

    title = (meta or {}).get("title") or job.get("title") or "New SOP"
    owner = m("owner", "[Content owner]")
    reviewer = m("reviewer", "[Content reviewer]")
    approver = m("approver", "[Content approver]")
    eff = m("effective_date", today.isoformat())
    cycle = m("review_cycle", "2 Years")
    try:
        years = int(str(cycle).split()[0])
        d0 = dt.date.fromisoformat(eff)
        next_review = d0.replace(year=d0.year + years).isoformat()
    except Exception:
        next_review = ""

    dept_cfg = template_cfg.get("departments", {}).get(job.get("department") or "", {})
    values = {
        "Title": title,
        "txt_vzi_SBU": m("sbu", "Gamsberg"),
        "txt_vzi_Dept": m("dept", dept_cfg.get("dept_label", "")),
        "VZI_DocNo": m("doc_no", "[Doc number]"),
        "SBU_DocNo": m("sbu_doc_no", ""),
        "txt_vzi_DocType": m("doc_type", "Standard Operating Procedure"),
        "txtStatus": "Draft",
        "ExtVersion": m("version", "0"),
        "txtSecurity": m("security", "Internal (C3)"),
        "txt_Audit_Standard": m("audit_standard", "ISO 45001:2018"),
        "txt_HOD": m("hod", "[HOD]"),
        "EffectiveDate": eff,
        "txtEffectiveDate": eff,
        "ReviewDate": eff,
        "txtReviewCycle": cycle,
        "NextReviewDate": next_review,
        "txtContentOwners": owner,
        "txtReviewers": reviewer,
        "txtApprovers": approver,
        "txtCW2": owner,
        "txtRew2": reviewer,
        "txtApp2": approver,
    }

    # content controls in body, headers and footers
    _fill_sdts(body, values)
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/header") or rel.reltype.endswith("/footer"):
            _fill_sdts(rel.target_part.element, values)
    doc.core_properties.title = title

    # ask Word to refresh the table of contents on open
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)

    elements = list(body)

    # ---- prototypes from template
    i_purpose = _body_index(elements, lambda e: _is_h1(e, "Purpose Scope"))
    i_valid = _body_index(elements, lambda e: _is_h1(e, "Validity and Document"))
    protos = {
        "h1": copy.deepcopy(elements[i_purpose]),
        "h2": copy.deepcopy(next(e for e in elements if e.tag == qn("w:p") and _text(e).strip() == "Document Validity")),
        "step": copy.deepcopy(next(e for e in elements if _text(e).startswith("This SOP replaces any other SOP"))),
    }
    for el in protos.values():
        for hl in list(el.iter(qn("w:highlight"))):
            hl.getparent().remove(hl)
    hazard_grid = copy.deepcopy(next(e for e in elements if e.tag == qn("w:tbl") and _text(e).startswith("Hazards / Energies")))

    # ---- change history: first data row
    ch = next(e for e in elements if e.tag == qn("w:tbl") and _text(e).startswith("Modified Date"))
    from docx.table import Table
    ch_t = Table(ch, doc)
    if len(ch_t.rows) > 1:
        for cell, val in zip(ch_t.rows[1].cells, [today.isoformat(), m("version", "0"), owner, "First draft issued for review (AI-assisted draft)"]):
            cell.text = val

    by_id = {s["id"]: s for s in sections}

    # ---- definitions (fixed slot)
    defs = next(e for e in elements if e.tag == qn("w:tbl") and _text(e).startswith("Word / Phrase"))
    i_defs_h = _body_index(elements, lambda e: _is_h1(e, "Definitions and Abbreviations"))
    if "definitions" in by_id:
        dt_t = Table(defs, doc)
        rows = []
        for b in by_id["definitions"]["blocks"]:
            if b["type"] == "table":
                rows += b["rows"]
        # drop template example rows, keep header
        for r in list(dt_t.rows)[1:]:
            r._tr.getparent().remove(r._tr)
        for r in rows:
            new = dt_t.add_row()
            new.cells[0].text = r[0] if r else ""
            new.cells[1].text = r[1] if len(r) > 1 else ""
    elif i_defs_h >= 0:
        body.remove(elements[i_defs_h])
        body.remove(defs)

    # ---- annexures (fixed slot at end)
    i_ann = _body_index(elements, lambda e: _is_h1(e, "Annexures"))
    sect_pr = body.find(qn("w:sectPr"))
    if i_ann >= 0:
        for el in elements[i_ann + 1:]:
            if el is not sect_pr:
                body.remove(el)
        if "annexures" in by_id:
            w = BodyWriter(doc, sect_pr, protos)
            for b in by_id["annexures"]["blocks"]:
                if b["type"] in ("bullets", "steps"):
                    for it in b["items"]:
                        w.bullet(it)
                else:
                    w.blocks([b])
        else:
            body.remove(elements[i_ann])

    # ---- main body: remove example sections, insert ours before "Validity and Document Management"
    anchor = elements[i_valid]
    for el in elements[i_purpose:i_valid]:
        body.remove(el)

    w = BodyWriter(doc, anchor, protos)
    for s in sections:
        if s["id"] in ("definitions", "annexures"):
            continue
        w.heading1(s["heading"])
        if s.get("verify"):
            w.para("Note: AI-assisted content. Verify every reference against the VZI legal register before approval.", color="C00000")
        w.blocks(s["blocks"])
        if s.get("hazard_grid"):
            w.spacer()
            _tick_hazards(hazard_grid, job.get("hazards") or [])
            w.raw(hazard_grid)
        w.spacer()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _tick_hazards(tbl, hazards: list[str]):
    chosen = {h.strip().lower() for h in hazards}
    for tr in tbl.findall(qn("w:tr")):
        tcs = tr.findall(qn("w:tc"))
        for i in range(0, len(tcs) - 1, 2):
            name = _text(tcs[i]).strip().lower()
            if name and name in chosen:
                p = tcs[i + 1].find(qn("w:p"))
                if p is not None:
                    _set_par_text(p, "✓", bold=True)
