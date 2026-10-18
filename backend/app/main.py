"""Agent 02 SOP drafting demo: FastAPI backend.

Run:  uvicorn app.main:app --reload   (from the backend folder)
"""
from __future__ import annotations

import io
import json
import os
import re
from pathlib import Path
from typing import Any, Optional

from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

ROOT = Path(__file__).resolve().parents[2]
load_dotenv(ROOT / ".env")

from . import llm  # noqa: E402  (after env is loaded)
from .render_docx import render  # noqa: E402

TEMPLATES_DIR = Path(__file__).resolve().parents[1] / "templates"
FRONTEND_DIST = ROOT / "frontend" / "dist"

app = FastAPI(title="VZI Agent 02 SOP Drafting Demo")


def load_templates() -> dict[str, dict]:
    out = {}
    for f in sorted(TEMPLATES_DIR.glob("*.json")):
        cfg = json.loads(f.read_text(encoding="utf-8"))
        out[cfg["id"]] = cfg
    return out


def get_template(template_id: str) -> dict:
    t = load_templates().get(template_id)
    if not t:
        raise HTTPException(404, f"Unknown template {template_id}")
    return t


# ------------------------------------------------------------------ models

class Section(BaseModel):
    id: str
    label: Optional[str] = None
    heading: str
    help: Optional[str] = None
    instruction: Optional[str] = None
    verify: Optional[bool] = None
    hazard_grid: Optional[bool] = None
    blocks: list[dict[str, Any]] = []


class DraftReq(BaseModel):
    template_id: str = "gamsberg_sop"
    job: dict[str, Any]
    section: Section
    all_titles: list[str] = []


class ReviseReq(BaseModel):
    template_id: str = "gamsberg_sop"
    job: dict[str, Any]
    section: Section
    instruction: str


class SuggestReq(BaseModel):
    template_id: str = "gamsberg_sop"
    job: dict[str, Any]
    kind: str


class ExportReq(BaseModel):
    template_id: str = "gamsberg_sop"
    job: dict[str, Any]
    meta: dict[str, Any] = {}
    sections: list[Section]


# ------------------------------------------------------------------ routes

@app.get("/api/health")
def health():
    return {"mode": "gemini" if llm.is_live() else "demo", "model": llm._model_name() if llm.is_live() else None}


@app.get("/api/templates")
def templates():
    return [{"id": t["id"], "name": t["name"], "site": t["site"]} for t in load_templates().values()]


@app.get("/api/templates/{template_id}")
def template(template_id: str):
    return get_template(template_id)


def _err(e: Exception):
    msg = str(e)
    if "API key" in msg or "API_KEY" in msg or "PERMISSION_DENIED" in msg:
        msg = "The Gemini API key was rejected. Check GEMINI_API_KEY in the .env file."
    elif "RESOURCE_EXHAUSTED" in msg or "429" in msg:
        msg = "The AI is busy (rate limit). Wait a few seconds and try again."
    return JSONResponse(status_code=502, content={"detail": msg})


@app.post("/api/draft-section")
def draft_section(req: DraftReq):
    t = get_template(req.template_id)
    try:
        blocks = llm.draft_section(t, req.job, req.section.model_dump(), req.all_titles)
    except Exception as e:
        return _err(e)
    return {"blocks": blocks}


@app.post("/api/revise-section")
def revise_section(req: ReviseReq):
    t = get_template(req.template_id)
    try:
        blocks = llm.revise_section(t, req.job, req.section.model_dump(), req.section.blocks, req.instruction)
    except Exception as e:
        return _err(e)
    return {"blocks": blocks}


@app.post("/api/suggest")
def suggest(req: SuggestReq):
    t = get_template(req.template_id)
    try:
        items = llm.suggest(t, req.job, req.kind)
    except Exception as e:
        return _err(e)
    return {"items": items}


@app.post("/api/extract")
async def extract(file: UploadFile = File(...)):
    """Read an uploaded SOP (PDF, DOCX or TXT) and return its text."""
    data = await file.read()
    name = (file.filename or "").lower()
    text = ""
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        elif name.endswith(".docx"):
            import docx

            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for tbl in d.tables:
                for row in tbl.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8", errors="ignore")
    except Exception as e:
        raise HTTPException(400, f"Could not read this file: {e}")
    # drop repeated page header/footer lines from controlled PDFs
    text = re.sub(r"Revisions of this document are controlled electronically.*\n?", "", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 50:
        raise HTTPException(400, "No readable text found. If this is a scanned PDF, please upload the Word version.")
    return {"filename": file.filename, "chars": len(text), "text": text}


@app.post("/api/export")
def export(req: ExportReq):
    t = get_template(req.template_id)
    data = render(TEMPLATES_DIR / t["docx"], t, req.job, req.meta, [s.model_dump() for s in req.sections])
    title = (req.meta.get("title") or req.job.get("title") or "SOP draft").strip()
    safe = re.sub(r"[^A-Za-z0-9 _-]+", "", title)[:80].strip() or "SOP draft"
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{safe} - DRAFT.docx"'},
    )


# ------------------------------------------------------------------ frontend

if FRONTEND_DIST.exists():
    app.mount("/assets", StaticFiles(directory=FRONTEND_DIST / "assets"), name="assets")

    @app.get("/{path:path}")
    def spa(path: str):
        f = FRONTEND_DIST / path
        if path and f.is_file():
            return FileResponse(f)
        return FileResponse(FRONTEND_DIST / "index.html")
